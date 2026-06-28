#!/usr/bin/env python3
"""
将 Steam Games XLSX 数据转换为标准 CSV 格式

用法:
    python convert_xlsx_to_csv.py

依赖: pip install openpyxl
"""

import csv
import sys

try:
    from openpyxl import load_workbook
except ImportError:
    print("[ERROR] 请安装 openpyxl: pip install openpyxl")
    sys.exit(1)

XLSX_FILE = "data/member1/games.csv"       # 实际是 XLSX 格式
OUTPUT_FILE = "data/member1/games_real.csv" # 输出真正的 CSV
SAMPLE_FILE = "data/member1/games-example.csv"
SAMPLE_OUT = "data/member1/games-example-real.csv"


def convert(input_path, output_path):
    print(f"读取: {input_path}")
    wb = load_workbook(input_path, read_only=True)
    ws = wb.active

    with open(output_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        row_count = 0
        for row in ws.iter_rows(values_only=True):
            writer.writerow(row)
            row_count += 1
            if row_count % 10000 == 0:
                print(f"  已转换 {row_count} 行...")

    print(f"完成: {output_path} ({row_count} 行)")
    wb.close()


if __name__ == "__main__":
    # 转完整数据
    convert(XLSX_FILE, OUTPUT_FILE)
    # 转样例数据
    convert(SAMPLE_FILE, SAMPLE_OUT)
    print("\n转换完成！现在用 games_real.csv 运行 Spark 分析")
